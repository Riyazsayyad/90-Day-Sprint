"""Merge exact LeetCode URLs from xlsx into assets/dsa-curriculum.json by Q#."""
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Install openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

XLSX = Path(r"c:\Riyaz\Risky-Downloads\dsa_curriculum_exact_leetcode_links.xlsx")
CURRICULUM = Path(__file__).resolve().parent.parent / "assets" / "dsa-curriculum.json"

URL_RE = re.compile(r"https://leetcode\.com/problems/[^\s|,)\"']+/?")


def extract_urls(cell) -> list[str]:
    if not cell:
        return []
    return URL_RE.findall(str(cell).strip())


def slug_from_url(url: str) -> str | None:
    m = re.search(r"/problems/([^/]+)/?", url)
    return m.group(1) if m else None


def load_link_map() -> dict[int, str]:
    if not XLSX.exists():
        print(f"XLSX not found: {XLSX}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(str(XLSX), read_only=True)
    ws = wb.active
    mapping: dict[int, str] = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        q = int(row[0])
        urls = extract_urls(row[5])
        if urls:
            mapping[q] = urls[0].rstrip("/") + "/"

    # Row 126 cell contains two problems (126 + 127); curriculum may lack q127.
    row126 = next(
        (r for r in ws.iter_rows(min_row=2, values_only=True) if str(r[0]) == "126"),
        None,
    )
    if row126:
        urls = extract_urls(row126[5])
        if len(urls) >= 2 and 127 not in mapping:
            mapping[127] = urls[1].rstrip("/") + "/"

    return mapping


def main():
    link_map = load_link_map()
    data = json.loads(CURRICULUM.read_text(encoding="utf-8"))

    linked = 0
    unlinked = 0
    for section in data["sections"]:
        for pattern in section["patterns"]:
            for problem in pattern["problems"]:
                url = link_map.get(problem["num"])
                if url:
                    problem["url"] = url
                    problem["leetcode"] = slug_from_url(url)
                    linked += 1
                else:
                    problem["url"] = None
                    problem["leetcode"] = None
                    unlinked += 1

    CURRICULUM.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Updated {CURRICULUM.name}: {linked} exact links, {unlinked} without link")


if __name__ == "__main__":
    main()
